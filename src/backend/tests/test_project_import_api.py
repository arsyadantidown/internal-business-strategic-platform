from datetime import date
from decimal import Decimal
from uuid import UUID, uuid4
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.api.project_import import usecase as import_usecase
from app.api.project_import.usecase import ProjectImportUseCase
from app.api.project_import.views import (
    download_project_import_template,
)
from app.api.projects.usecase import CODE_TAKEN
from app.lib.portfolio_import import build_portfolio_import_template
from app.main import app
from app.models.business_unit_model import BusinessUnit
from app.models.company_size_model import CompanySize
from app.models.customer_type_model import CustomerType
from app.models.industry_model import Industry
from app.models.legal_structure_model import LegalStructure
from app.models.project_status_model import ProjectStatus
from app.models.tech_stack_model import TechStack
from app.models.customer_model import Customer
from app.models.partner_model import Partner
from app.models.project_financial_model import ProjectFinancial
from app.models.project_model import Project
from app.models.project_partner_model import ProjectPartner
from app.models.project_product_model import ProjectProduct
from app.models.project_tech_stack_model import ProjectTechStack
from tests.test_projects import project_payload


def test_project_import_template_route_registered():
    matching_routes = [
        route
        for route in app.routes
        if route.path
        == "/api/v1/projects/import/template"
    ]

    assert len(matching_routes) == 1

    route = matching_routes[0]

    assert "GET" in route.methods


def test_download_project_import_template_response():
    use_case = ProjectImportUseCase(
        db=None,
    )

    response = download_project_import_template(
        _portfolio_user=None,
        use_case=use_case,
    )

    assert response.status_code == 200

    assert response.media_type == (
        "application/vnd.openxmlformats-"
        "officedocument.spreadsheetml.sheet"
    )

    assert response.headers[
        "content-disposition"
    ] == (
        'attachment; filename="'
        'portfolio_import_template.xlsx"'
    )

    workbook = load_workbook(
        BytesIO(response.body),
        read_only=False,
    )

    worksheet = workbook["Sheet1"]

    assert (
        worksheet["A1"].value
        == "Customer Perspective"
    )
    assert (
        worksheet["I1"].value
        == "Project Perspective"
    )
    assert (
        worksheet["AC1"].value
        == "Financial Perspective"
    )
    assert (
        worksheet["AG1"].value
        == "Tech Stack"
    )
    assert (
        worksheet["AK1"].value
        == "Partner Perspective"
    )

    assert worksheet["A2"].value == "No"
    assert worksheet["B2"].value == "Customer"
    assert worksheet["O2"].value == "Project ID"
    assert (
        worksheet["P2"].value
        == "Project Value (before tax)"
    )
    assert worksheet["AK2"].value == "Partner"

    assert worksheet.max_column == 37

    workbook.close()

PREVIEW_URL = "/api/v1/projects/import/preview"
CONFIRM_URL = "/api/v1/projects/import/confirm"
PROJECTS_URL = "/api/v1/projects"


async def create_existing_project(client, admin, refs):
    response = await client.post(
        PROJECTS_URL,
        headers=admin,
        json=project_payload(refs, code="IMPORT_EXISTING"),
    )
    assert response.status_code == 201, response.text
    return response.json()["data"]["id"]


def import_workbook(rows):
    workbook = load_workbook(BytesIO(build_portfolio_import_template()))
    worksheet = workbook["Sheet1"]
    headers = [cell.value for cell in worksheet[2]]
    for overrides in rows:
        values = {
            "Customer": "Preview Customer",
            "Project Name": "Preview Project",
            "Year": 2026,
            "Business Unit": "ITMS",
            "Project ID": "IMPORT_NEW",
            "Project Value (before tax)": 123,
            "Project Status": "On Going",
            "Start Date": date(2026, 1, 1),
            "End Date": date(2026, 12, 31),
        }
        values.update(overrides)
        worksheet.append([values.get(header) for header in headers])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    return output.getvalue()


async def upload_import(client, headers, content, *, url=CONFIRM_URL, filename="import.xlsx"):
    return await client.post(
        url,
        headers=headers,
        files={
            "file": (
                filename,
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


async def preview_rows(client, admin, rows):
    response = await upload_import(client, admin, import_workbook(rows), url=PREVIEW_URL)
    assert response.status_code == 200, response.text
    assert response.json()["success"] is True
    return response.json()["data"]


@pytest.mark.parametrize("project_code", ["IMPORT_EXISTING", "  import_existing  "])
async def test_preview_rejects_existing_project_code(client, admin, refs, project_code):
    await create_existing_project(client, admin, refs)
    data = await preview_rows(client, admin, [{"Project ID": project_code}])

    assert data["total_rows"] == 1
    assert data["valid_rows"] == 0
    assert data["invalid_rows"] == 1
    row = data["rows"][0]
    assert row["valid"] is False
    assert row["errors"] == [
        {
            "field": "project_code",
            "code": "PROJECT_CODE_EXISTS",
            "message": CODE_TAKEN,
        }
    ]


async def test_preview_preserves_errors_and_accepts_unseen_code(client, admin, refs):
    await create_existing_project(client, admin, refs)
    data = await preview_rows(
        client,
        admin,
        [
            {"Project ID": "IMPORT_EXISTING", "Project Name": None, "Business Unit": "UNKNOWN_UNIT"},
            {"Project ID": "import_existing"},
            {"Project ID": "IMPORT_NEW"},
        ],
    )

    assert data["total_rows"] == 3
    assert data["valid_rows"] == 1
    assert data["invalid_rows"] == 2
    for row in data["rows"][:2]:
        assert row["valid"] is False
        codes = [error["code"] for error in row["errors"]]
        assert "DUPLICATE_PROJECT_CODE" in codes
        assert codes.count("PROJECT_CODE_EXISTS") == 1
    errors = data["rows"][0]["errors"]
    assert any(error["field"] == "project_name" and error["code"] == "REQUIRED_FIELD" for error in errors)
    assert any(error["field"] == "business_unit" and error["code"] == "MASTER_DATA_NOT_FOUND" for error in errors)
    assert data["rows"][2]["valid"] is True
    assert data["rows"][2]["errors"] == []


async def test_preview_does_not_create_or_update_projects(client, admin, refs, db):
    await create_existing_project(client, admin, refs)
    query = select(Project.__table__).order_by(Project.id)
    before = [dict(row) for row in (await db.execute(query)).mappings().all()]
    assert len(before) == 1

    data = await preview_rows(
        client,
        admin,
        [{"Project ID": "IMPORT_EXISTING"}, {"Project ID": "IMPORT_NEW"}],
    )
    assert data["valid_rows"] == 1
    assert data["invalid_rows"] == 1
    after = [dict(row) for row in (await db.execute(query)).mappings().all()]
    assert after == before


async def test_preview_allows_code_of_soft_deleted_project(client, admin, refs):
    project_id = await create_existing_project(client, admin, refs)
    response = await client.delete(f"{PROJECTS_URL}/{project_id}", headers=admin)
    assert response.status_code == 200, response.text

    data = await preview_rows(client, admin, [{"Project ID": "import_existing"}])
    assert data["valid_rows"] == 1
    assert data["invalid_rows"] == 0
    assert data["rows"][0]["valid"] is True
    assert data["rows"][0]["errors"] == []


async def import_record_counts(db):
    counts = {}
    for model in (
        Project, Customer, ProjectFinancial, ProjectProduct,
        ProjectTechStack, ProjectPartner, Partner,
    ):
        counts[model.__tablename__] = (
            await db.execute(select(func.count()).select_from(model))
        ).scalar_one()
    return counts


async def test_confirm_imports_multiple_projects(client, admin, refs):
    content = import_workbook([
        {
            "Project ID": "CONFIRM_ONE",
            "Customer": "  Import Customer  ",
            "Customer Size": "Tier 1",
            "Customer Type": "Strategic Customer",
            "Industry": "Government",
            "Legal Structure": "State-owned Enterprise",
            "Location": "Bandung",
            "Project Value (before tax)": 1000,
            "Carry Over Value": 50,
            "COGS": 200,
            "Operational Costs": 100,
            "Product Use": "Bliv Package",
            "Frontend": "React",
            "Backend": "Golang",
            "Data & Storage": "Hadoop",
            "Contract/PO Number": "IMPORT-CONTRACT",
        },
        {"Project ID": "CONFIRM_TWO"},
    ])
    response = await upload_import(client, admin, content)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["success"] is True
    assert body["message"] == "Import project berhasil disimpan."
    data = body["data"]
    assert data["imported_count"] == 2
    assert [item["project_code"] for item in data["projects"]] == ["CONFIRM_ONE", "CONFIRM_TWO"]
    assert all(set(item) == {"id", "project_code"} for item in data["projects"])
    assert len({item["id"] for item in data["projects"]}) == 2

    detail = await client.get(f"{PROJECTS_URL}/{data['projects'][0]['id']}", headers=admin)
    assert detail.status_code == 200, detail.text
    project = detail.json()["data"]
    assert project["customer"]["customer_name"] == "Import Customer"
    assert project["customer"]["customer_location"] == "Bandung"
    assert project["customer"]["company_size"]["id"] == refs["company_size"]["TIER_1"]
    assert project["customer"]["industry"]["id"] == refs["industry"]["GOVERNMENT"]
    assert project["contract_number"] == "IMPORT-CONTRACT"
    assert Decimal(str(project["carry_over_value"])) == Decimal("50")
    assert Decimal(str(project["financial"]["net_profit"])) == Decimal("700")
    for category, name in (
        ("products", "Bliv Package"), ("frontend", "React"),
        ("backend", "Golang"), ("data_storage", "Hadoop"),
    ):
        assert [item["name"] for item in project["tech_stack"][category]] == [name]

    detail = await client.get(f"{PROJECTS_URL}/{data['projects'][1]['id']}", headers=admin)
    assert detail.status_code == 200, detail.text
    financial = detail.json()["data"]["financial"]
    assert Decimal(str(financial["cogs"])) == 0
    assert Decimal(str(financial["operational_cost"])) == 0


async def test_confirm_revalidates_code_after_preview(client, admin, refs, db):
    content = import_workbook([{"Project ID": "IMPORT_EXISTING"}])
    preview = await upload_import(client, admin, content, url=PREVIEW_URL)
    assert preview.status_code == 200, preview.text
    assert preview.json()["data"]["valid_rows"] == 1
    await create_existing_project(client, admin, refs)
    before = await import_record_counts(db)

    response = await upload_import(client, admin, content)
    assert response.status_code == 422, response.text
    row = response.json()["error"]["details"]["rows"][0]
    assert row["row"] == 3
    assert any(error["code"] == "PROJECT_CODE_EXISTS" for error in row["errors"])
    assert await import_record_counts(db) == before


@pytest.mark.parametrize("invalid", [{"Project Name": None}, {"Year": 1800}])
async def test_confirm_invalid_row_saves_nothing(client, admin, db, invalid):
    before = await import_record_counts(db)
    response = await upload_import(client, admin, import_workbook([
        {"Project ID": "VALID_ROW", "Partner": f"Vendor: Pending {uuid4().hex}"},
        {"Project ID": "INVALID_ROW", **invalid},
    ]))
    assert response.status_code == 422, response.text
    assert response.json()["error"]["code"] == "IMPORT_VALIDATION_ERROR"
    assert response.json()["error"]["details"]["rows"][0]["row"] == 4
    assert await import_record_counts(db) == before


@pytest.mark.parametrize("failure", ["http", "unexpected"])
async def test_confirm_rolls_back_projects_and_new_partners(
    client, admin, db, monkeypatch, failure
):
    before = await import_record_counts(db)
    original_create = import_usecase.create_project
    calls = []

    async def fail_second(session, payload, *, created_by):
        calls.append(payload.project_code)
        if len(calls) == 2:
            assert (await session.execute(select(func.count(Project.id)))).scalar_one() == 1
            assert (await session.execute(select(func.count(Partner.id)))).scalar_one() == before["partners"] + 1
            if failure == "http":
                raise HTTPException(409, "Simulated creation conflict")
            raise RuntimeError("Simulated creation failure")
        return await original_create(session, payload, created_by=created_by)

    monkeypatch.setattr(import_usecase, "create_project", fail_second)
    name = f"Rollback Partner {uuid4().hex}"
    content = import_workbook([
        {"Project ID": "ROLLBACK_ONE", "Partner": f"Principal: {name}"},
        {"Project ID": "ROLLBACK_TWO", "Partner": f"Vendor: {name}"},
    ])
    if failure == "http":
        response = await upload_import(client, admin, content)
        assert response.status_code == 409, response.text
        assert response.json()["error"]["details"]["rows"][0]["row"] == 4
    else:
        with pytest.raises(RuntimeError, match="Simulated creation failure"):
            await upload_import(client, admin, content)
    assert calls == ["ROLLBACK_ONE", "ROLLBACK_TWO"]
    assert await import_record_counts(db) == before


async def test_confirm_reuses_existing_partner(client, admin, refs, db):
    partner_id = UUID(refs["partners"]["alpha"])
    partner = await db.get(Partner, partner_id)
    before = (await db.execute(select(func.count(Partner.id)))).scalar_one()
    response = await upload_import(client, admin, import_workbook([
        {"Project ID": "PARTNER_EXISTING", "Partner": f"Principal: {partner.partner_name.upper()}"},
    ]))
    assert response.status_code == 200, response.text
    assert (await db.execute(select(func.count(Partner.id)))).scalar_one() == before
    link = (await db.execute(select(ProjectPartner))).scalar_one()
    assert link.partner_id == partner_id
    assert link.partner_role == "principal"


async def test_confirm_creates_shared_partner_once(client, admin, db):
    before = (await db.execute(select(func.count(Partner.id)))).scalar_one()
    name = f"Shared Partner {uuid4().hex}"
    response = await upload_import(client, admin, import_workbook([
        {"Project ID": "SHARED_ONE", "Partner": f"Principal: {name}"},
        {"Project ID": "SHARED_TWO", "Partner": f"Vendor:   {name.upper()}  "},
    ]))
    assert response.status_code == 200, response.text
    assert (await db.execute(select(func.count(Partner.id)))).scalar_one() == before + 1
    links = (await db.execute(select(ProjectPartner))).scalars().all()
    assert len(links) == 2
    assert len({link.partner_id for link in links}) == 1
    assert {link.partner_role for link in links} == {"principal", "vendor"}


@pytest.mark.parametrize("user_fixture", ["commercial", "engineer"])
async def test_confirm_requires_both_permissions(client, commercial, engineer, user_fixture, db):
    headers = {"commercial": commercial, "engineer": engineer}[user_fixture]
    before = await import_record_counts(db)
    response = await upload_import(client, headers, import_workbook([{}]))
    assert response.status_code == 403, response.text
    assert await import_record_counts(db) == before


@pytest.mark.parametrize(
    "filename, content",
    [("import.csv", b"not excel"), ("import.xlsx", b"not excel"), ("import.xlsx", b"")],
)
async def test_confirm_rejects_invalid_files_like_preview(client, admin, db, filename, content):
    before = await import_record_counts(db)
    preview = await upload_import(client, admin, content, url=PREVIEW_URL, filename=filename)
    response = await upload_import(client, admin, content, filename=filename)
    assert preview.status_code == response.status_code == 422
    assert preview.json()["message"] == response.json()["message"]
    assert await import_record_counts(db) == before


async def test_confirm_rejects_empty_workbook(client, admin, db):
    before = await import_record_counts(db)
    response = await upload_import(client, admin, import_workbook([]))
    assert response.status_code == 422, response.text
    assert await import_record_counts(db) == before


async def test_preview_with_new_partner_remains_read_only(client, admin, db):
    before = await import_record_counts(db)
    data = await preview_rows(client, admin, [
        {"Partner": f"Vendor: Preview Only {uuid4().hex}"},
    ])
    assert data["valid_rows"] == 1
    assert await import_record_counts(db) == before


async def create_customer_projects(client, admin, refs, names):
    for index, name in enumerate(names):
        response = await client.post(
            PROJECTS_URL,
            headers=admin,
            json=project_payload(refs, code=f"CUSTOMER_SOURCE_{index}", customer_name=name),
        )
        assert response.status_code == 201, response.text


async def test_similar_customer_warning_is_read_only(client, admin, refs, db):
    await create_customer_projects(client, admin, refs, ["PT  Bank\nMandiri"])
    before_counts = await import_record_counts(db)
    customer_query = select(Customer.__table__).order_by(Customer.id)
    project_query = select(Project.__table__).order_by(Project.id)
    before_customers = [dict(row) for row in (await db.execute(customer_query)).mappings().all()]
    before_projects = [dict(row) for row in (await db.execute(project_query)).mappings().all()]

    data = await preview_rows(client, admin, [{"Customer": "  pt BANK  mandirii  "}])
    assert data["valid_rows"] == 1
    assert data["invalid_rows"] == 0
    row = data["rows"][0]
    assert row["valid"] is True
    assert row["errors"] == []
    assert row["warnings"] == [{
        "field": "customer_name",
        "code": "SIMILAR_CUSTOMER_NAME",
        "message": "Nama customer mirip dengan customer yang sudah ada: PT Bank Mandiri",
    }]
    assert await import_record_counts(db) == before_counts
    assert [dict(row) for row in (await db.execute(customer_query)).mappings().all()] == before_customers
    assert [dict(row) for row in (await db.execute(project_query)).mappings().all()] == before_projects


async def test_customer_warning_does_not_block_confirm(client, admin, refs):
    await create_customer_projects(client, admin, refs, ["PT Bank Mandiri"])
    content = import_workbook([{"Customer": "PT Bank Mandirii"}])
    preview = await upload_import(client, admin, content, url=PREVIEW_URL)
    assert preview.status_code == 200, preview.text
    assert preview.json()["data"]["rows"][0]["warnings"][0]["code"] == "SIMILAR_CUSTOMER_NAME"

    response = await upload_import(client, admin, content)
    assert response.status_code == 200, response.text
    assert response.json()["data"]["imported_count"] == 1
    project_id = response.json()["data"]["projects"][0]["id"]
    detail = await client.get(f"{PROJECTS_URL}/{project_id}", headers=admin)
    assert detail.status_code == 200, detail.text
    assert detail.json()["data"]["customer"]["customer_name"] == "PT Bank Mandirii"


async def test_exact_normalized_customer_suppresses_warning(client, admin, refs):
    # A nearby alternative must not produce a warning when an exact match exists.
    await create_customer_projects(
        client, admin, refs, ["PT  Bank\nMandiri", "PT Bank Mandirix"]
    )
    data = await preview_rows(client, admin, [{"Customer": "  pt BANK\tmandiri  "}])
    assert data["rows"][0]["valid"] is True
    assert data["rows"][0]["warnings"] == []


@pytest.mark.parametrize(
    "existing, imported",
    [
        ("PT Bank Mandiri", "Unrelated Agricultural Cooperative"),
        # This close-looking pair scores about 0.903, below the 0.92 cutoff.
        ("PT Bank Mandira", "PT Bank Mandirii"),
    ],
)
async def test_customer_below_similarity_threshold_has_no_warning(
    client, admin, refs, existing, imported
):
    await create_customer_projects(client, admin, refs, [existing])
    data = await preview_rows(client, admin, [{"Customer": imported}])
    assert data["rows"][0]["valid"] is True
    assert data["rows"][0]["warnings"] == []


async def test_customer_warning_returns_only_best_match(client, admin, refs):
    await create_customer_projects(
        client, admin, refs, ["PT Bank Mandirix", "PT Bank Mandiri", "PT Bank Mandiri"]
    )
    data = await preview_rows(client, admin, [{"Customer": "PT Bank Mandirii"}])
    warnings = data["rows"][0]["warnings"]
    assert len(warnings) == 1
    assert warnings[0]["code"] == "SIMILAR_CUSTOMER_NAME"
    assert warnings[0]["message"].endswith(": PT Bank Mandiri")


async def test_customer_warning_breaks_equal_scores_deterministically(client, admin, refs):
    # Insert in reverse alphabetical order; both candidates score equally.
    await create_customer_projects(client, admin, refs, ["PT Bank Mandiro", "PT Bank Mandira"])
    data = await preview_rows(client, admin, [{"Customer": "PT Bank Mandiri"}])
    warnings = data["rows"][0]["warnings"]
    assert len(warnings) == 1
    assert warnings[0]["message"].endswith(": PT Bank Mandira")


async def test_customer_warning_preserves_existing_errors(client, admin, refs):
    await create_customer_projects(client, admin, refs, ["PT Bank Mandiri"])
    data = await preview_rows(client, admin, [{
        "Project ID": "CUSTOMER_SOURCE_0",
        "Project Name": None,
        "Customer": "PT Bank Mandirii",
    }])
    row = data["rows"][0]
    assert data["valid_rows"] == 0
    assert data["invalid_rows"] == 1
    assert row["valid"] is False
    assert {error["code"] for error in row["errors"]} == {"REQUIRED_FIELD", "PROJECT_CODE_EXISTS"}
    assert len(row["warnings"]) == 1
    assert row["warnings"][0]["code"] == "SIMILAR_CUSTOMER_NAME"


async def assert_invalid_preview_and_confirm(client, admin, db, overrides, field):
    before = await import_record_counts(db)
    content = import_workbook([
        {"Project ID": "VALID_VALIDATION_ROW", "Partner": f"Vendor: Pending {uuid4().hex}"},
        {"Project ID": "INVALID_VALIDATION_ROW", **overrides},
    ])
    response = await upload_import(client, admin, content, url=PREVIEW_URL)
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert (data["total_rows"], data["valid_rows"], data["invalid_rows"]) == (2, 1, 1)
    assert data["rows"][0]["valid"] is True
    row = data["rows"][1]
    assert row["row"] == 4
    assert row["valid"] is False
    assert any(error["field"] == field for error in row["errors"])
    assert len(row["errors"]) == len({
        (error["field"], error["code"], error["message"]) for error in row["errors"]
    })
    assert await import_record_counts(db) == before

    response = await upload_import(client, admin, content)
    assert response.status_code == 422, response.text
    assert response.json()["error"]["code"] == "IMPORT_VALIDATION_ERROR"
    assert response.json()["error"]["details"]["rows"] == [
        {"row": 4, "errors": row["errors"]}
    ]
    assert await import_record_counts(db) == before
    return row


@pytest.mark.parametrize(
    "header,value,field",
    [
        ("Year", 1899, "year"),
        ("Year", 2201, "year"),
        ("COGS", -1, "cogs"),
        ("Operational Costs", -1, "operational_cost"),
        ("Carry Over Value", -1, "carry_over_value"),
        ("Project Value (before tax)", -1, "value_before_tax"),
        ("COGS", "abc", "cogs"),
        ("Operational Costs", "abc", "operational_cost"),
        ("Carry Over Value", "abc", "carry_over_value"),
        ("Project Value (before tax)", "abc", "value_before_tax"),
        ("COGS", "NaN", "cogs"),
        ("Operational Costs", "Infinity", "operational_cost"),
        ("COGS", "1.001", "cogs"),
        ("Project Value (before tax)", "123456789012345678901", "value_before_tax"),
        ("Project ID", "P" * 101, "project_code"),
        ("Project Name", "P" * 256, "project_name"),
        ("Customer", "C" * 256, "customer_name"),
        ("Location", "L" * 256, "customer_location"),
        ("Revenue/Annual Sales/Annual Budget", "R" * 101, "annual_revenue"),
        ("Project Manager", "M" * 151, "project_manager"),
        ("Customer Size", 1.5, "customer_size"),
        ("Customer Size", 7, "customer_size_code"),
    ],
)
async def test_preview_creation_constraints_reject_without_saving(
    client, admin, db, header, value, field
):
    await assert_invalid_preview_and_confirm(client, admin, db, {header: value}, field)


async def test_preview_preserves_existing_errors_with_schema_errors(client, admin, db):
    row = await assert_invalid_preview_and_confirm(
        client, admin, db,
        {"Project Name": None, "Year": 1899, "COGS": "abc"}, "year",
    )
    assert {error["field"] for error in row["errors"]} == {"project_name", "year", "cogs"}
    assert len(row["errors"]) == 3
    assert next(error for error in row["errors"] if error["field"] == "project_name")["code"] == "REQUIRED_FIELD"


@pytest.mark.parametrize("blank", [None, "   "])
async def test_preview_blank_optional_money_stays_valid(client, admin, db, blank):
    before = await import_record_counts(db)
    content = import_workbook([{
        "COGS": blank, "Operational Costs": blank, "Carry Over Value": blank,
    }])
    response = await upload_import(client, admin, content, url=PREVIEW_URL)
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert (data["valid_rows"], data["invalid_rows"]) == (1, 0)
    assert data["rows"][0]["errors"] == []
    assert await import_record_counts(db) == before

    response = await upload_import(client, admin, content)
    assert response.status_code == 200, response.text
    financial = (await db.execute(select(ProjectFinancial))).scalar_one()
    assert financial.cogs == financial.operational_cost == Decimal("0")
    project = (await db.execute(select(Project))).scalar_one()
    assert project.carry_over_value is None


@pytest.mark.parametrize("tier", [1, 1.0, "1", "1.0", "Tier 1"])
async def test_preview_numeric_and_text_tiers_resolve(client, admin, refs, db, tier):
    before = await import_record_counts(db)
    data = await preview_rows(client, admin, [{"Customer Size": tier}])
    assert (data["valid_rows"], data["invalid_rows"]) == (1, 0)
    assert data["rows"][0]["errors"] == []
    assert data["rows"][0]["data"]["customer_size_code"] == "TIER_1"
    assert data["rows"][0]["data"]["company_size_id"] == refs["company_size"]["TIER_1"]
    assert await import_record_counts(db) == before


@pytest.mark.parametrize("year", [1900, 2200])
async def test_preview_accepts_schema_year_boundaries(client, admin, year):
    data = await preview_rows(client, admin, [{"Year": year}])
    assert (data["valid_rows"], data["invalid_rows"]) == (1, 0)
    assert data["rows"][0]["errors"] == []


@pytest.mark.parametrize(
    "model,group,code,header,value,field",
    [
        (BusinessUnit, "business_unit", "CSD", "Business Unit", "CSD", "business_unit_id"),
        (ProjectStatus, "status", "CLOSED", "Project Status", "Closed", "project_status_id"),
        (CompanySize, "company_size", "TIER_1", "Customer Size", 1, "company_size_id"),
        (CustomerType, "customer_type", "STRATEGIC_CUSTOMER", "Customer Type", "Strategic Customer", "customer_type_id"),
        (Industry, "industry", "GOVERNMENT", "Industry", "Government", "industry_id"),
        (LegalStructure, "legal_structure", "STATE_OWNED_ENTERPRISE", "Legal Structure", "State-owned Enterprise", "legal_structure_id"),
    ],
)
async def test_preview_rejects_inactive_master_reference(
    client, admin, refs, db, model, group, code, header, value, field
):
    item = await db.get(model, UUID(refs[group][code]))
    original_active = item.is_active
    try:
        item.is_active = False
        await db.commit()
        row = await assert_invalid_preview_and_confirm(
            client, admin, db, {header: value}, field,
        )
        assert any(error["code"] == "INVALID_REFERENCE" for error in row["errors"])
    finally:
        await db.rollback()
        await db.refresh(item)
        item.is_active = original_active
        await db.commit()


async def test_preview_rejects_inactive_tech_stack(client, admin, refs, db):
    item = await db.get(TechStack, UUID(refs["tech"]["frontend"][0]))
    original_active = item.is_active
    value = '"' + item.tech_stack_name.replace('"', '""') + '"'
    try:
        item.is_active = False
        await db.commit()
        row = await assert_invalid_preview_and_confirm(
            client, admin, db, {"Frontend": value}, "tech_stack_ids",
        )
        assert any(error["code"] == "INVALID_REFERENCE" for error in row["errors"])
    finally:
        await db.rollback()
        await db.refresh(item)
        item.is_active = original_active
        await db.commit()
