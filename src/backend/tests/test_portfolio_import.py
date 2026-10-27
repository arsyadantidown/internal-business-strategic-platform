from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.lib.portfolio_import import (
    extract_customer_size_code,
    normalize_decimal,
    normalize_text,
    parse_partner_cell,
    parse_portfolio_row,
    read_portfolio_workbook,
    split_multi_value,
    validate_import_rows,
    validate_portfolio_row,
    build_portfolio_import_template,
)


def test_normalize_text():
    result = normalize_text(
        "  PT Bangunindo \n Teknusa   Jaya "
    )

    assert result == "PT Bangunindo Teknusa Jaya"


def test_extract_customer_size_code():
    result = extract_customer_size_code(
        "Tier 1 : Strategic/Enterprise (>1T)"
    )

    assert result == "TIER_1"


def test_split_multi_value_with_quoted_comma():
    result = split_multi_value(
        '"Vue.js (Vue, Vue2, Vue3)", Nuxt 3 (JavaScript)'
    )

    assert result == [
        "Vue.js (Vue, Vue2, Vue3)",
        "Nuxt 3 (JavaScript)",
    ]


def test_parse_partner_cell():
    result = parse_partner_cell(
        "Principal: Zabbix SIA, \n"
        "Vendor: PT Bangunindo Teknusa"
    )

    assert result == [
        {
            "partner_role": "principal",
            "partner_name": "Zabbix SIA",
        },
        {
            "partner_role": "vendor",
            "partner_name": "PT Bangunindo Teknusa",
        },
    ]


def test_normalize_decimal():
    assert normalize_decimal(
        3500000000.0
    ) == Decimal("3500000000.0")

    assert normalize_decimal(None) is None


def test_parse_portfolio_row():
    headers = [
        "No",
        "Customer",
        "Customer Size",
        "Revenue/Annual Sales/Annual Budget",
        "Customer Type",
        "Industry",
        "Location",
        "Legal Structure",
        "Opportunity",
        "Project Name",
        "Year",
        "Contract/PO Name",
        "Contract/PO Number",
        "Business Unit",
        "Project ID",
        "Project Value (before tax)",
        "Project Value (after tax)",
        "Carry Over Value",
        "Project Status",
        "Progress Completion ",
        "%",
        "Project Manager",
        "Account Manager",
        "Start Date",
        "End Date",
        "Total Days",
        "Remaining Days",
        "Project Objective",
        "Revenue",
        "COGS",
        "Operational Costs",
        "Net Profit",
        "Product Use",
        "Frontend",
        "Backend",
        "Data & Storage",
        "Partner",
    ]

    row = [
        1,
        "PT Bank Mandiri (Persero) Tbk",
        "Tier 1 : Strategic/Enterprise (>1T)",
        "115T (FY 2025)",
        "Strategic Customer",
        "Financial Services",
        "Jakarta",
        "State-owned Enterprise",
        "Bank Mandiri SIEM & Monitoring",
        "Bank Mandiri SIEM & Monitoring",
        2026.0,
        "Enterprise Infrastructure Monitoring 2026",
        "PO-2026-BM-01824 / 4500091823",
        "ITMS",
        "ITMS_MON_BMRI_2026_PRO",
        3500000000.0,
        3885000000.0,
        None,
        "On Going",
        None,
        50,
        "Ahmad Fauzi",
        "Rizky Pratama",
        datetime(2026, 9, 1),
        datetime(2027, 8, 31),
        364,
        300,
        (
            "Implementasi SIEM & monitoring proaktif untuk "
            "keandalan dan keamanan infrastruktur IT enterprise"
        ),
        3500000000.0,
        2100000000.0,
        350000000.0,
        1050000000.0,
        "Bliv Package, Bliv Dashboard, Bliv Pipeline",
        '"Vue.js (Vue, Vue2, Vue3)", Nuxt 3 (JavaScript)',
        '"Python (Flask API, FastAPI)", "PHP (CI3, CI4)"',
        "PostgreSQL (Postgres), Hadoop",
        (
            "Principal: Zabbix SIA, \n"
            "Vendor: PT Bangunindo Teknusa"
        ),
    ]

    result = parse_portfolio_row(
        headers=headers,
        row=row,
        row_number=3,
    )

    assert result["source_row"] == 3

    assert result["customer_name"] == (
        "PT Bank Mandiri (Persero) Tbk"
    )

    assert result["customer_size_code"] == "TIER_1"

    assert result["project_name"] == (
        "Bank Mandiri SIEM & Monitoring"
    )

    assert result["project_code"] == (
        "ITMS_MON_BMRI_2026_PRO"
    )

    assert result["year"] == 2026

    assert result["start_date"] == date(2026, 9, 1)
    assert result["end_date"] == date(2027, 8, 31)

    assert result["value_before_tax"] == Decimal(
        "3500000000.0"
    )

    assert result["cogs"] == Decimal(
        "2100000000.0"
    )

    assert result["operational_cost"] == Decimal(
        "350000000.0"
    )

    assert result["products"] == [
        "Bliv Package",
        "Bliv Dashboard",
        "Bliv Pipeline",
    ]

    assert result["frontend"] == [
        "Vue.js (Vue, Vue2, Vue3)",
        "Nuxt 3 (JavaScript)",
    ]

    assert result["backend"] == [
        "Python (Flask API, FastAPI)",
        "PHP (CI3, CI4)",
    ]

    assert result["partners"] == [
        {
            "partner_role": "principal",
            "partner_name": "Zabbix SIA",
        },
        {
            "partner_role": "vendor",
            "partner_name": "PT Bangunindo Teknusa",
        },
    ]

    # Field hasil formula Excel tidak boleh digunakan.
    assert "value_after_tax" not in result
    assert "revenue" not in result
    assert "net_profit" not in result
    assert "total_days" not in result
    assert "remaining_days" not in result
    assert "progress_percent" not in result


def make_valid_row(
    project_code="PRJ-001",
    source_row=3,
):
    """
    Helper untuk membuat row valid.
    """
    return {
        "customer_name": "PT Test Customer",
        "project_name": "Test Project",
        "year": 2026,
        "business_unit": "ITMS",
        "project_code": project_code,
        "value_before_tax": Decimal("100000000"),
        "project_status": "On Going",
        "start_date": date(2026, 1, 1),
        "end_date": date(2026, 12, 31),
        "source_row": source_row,
    }


def test_validate_valid_portfolio_row():
    row = make_valid_row()

    errors = validate_portfolio_row(row)

    assert errors == []


def test_validate_missing_required_field():
    row = make_valid_row()

    row["project_name"] = None

    errors = validate_portfolio_row(row)

    assert any(
        error["field"] == "project_name"
        and error["code"] == "REQUIRED_FIELD"
        for error in errors
    )


def test_validate_invalid_date_range():
    row = make_valid_row()

    row["start_date"] = date(2026, 12, 31)
    row["end_date"] = date(2026, 1, 1)

    errors = validate_portfolio_row(row)

    assert any(
        error["code"] == "INVALID_DATE_RANGE"
        for error in errors
    )


def test_validate_negative_project_value():
    row = make_valid_row()

    row["value_before_tax"] = Decimal("-1")

    errors = validate_portfolio_row(row)

    assert any(
        error["field"] == "value_before_tax"
        and error["code"] == "INVALID_VALUE"
        for error in errors
    )


def test_duplicate_project_code_in_file():
    rows = [
        make_valid_row(
            project_code="PRJ-001",
            source_row=3,
        ),
        make_valid_row(
            project_code="prj-001",
            source_row=4,
        ),
    ]

    result = validate_import_rows(rows)

    assert result["total_rows"] == 2
    assert result["valid_rows"] == 0
    assert result["invalid_rows"] == 2

    assert result["rows"][0]["valid"] is False
    assert result["rows"][1]["valid"] is False

    assert any(
        error["code"] == "DUPLICATE_PROJECT_CODE"
        for error in result["rows"][0]["errors"]
    )

    assert any(
        error["code"] == "DUPLICATE_PROJECT_CODE"
        for error in result["rows"][1]["errors"]
    )


def test_read_portfolio_workbook():
    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Sheet1"

    # Row 1 = nama perspective.
    worksheet.append(
        [
            "Customer Perspective",
            "Project Perspective",
        ]
    )

    # Row 2 = header.
    worksheet.append(
        [
            "Customer",
            "Project Name",
            "Year",
            "Business Unit",
            "Project ID",
            "Project Value (before tax)",
            "Project Status",
            "Start Date",
            "End Date",
        ]
    )

    # Row 3 = data valid.
    worksheet.append(
        [
            "PT Test Customer",
            "Project Test",
            2026,
            "ITMS",
            "PRJ-001",
            100000000,
            "On Going",
            datetime(2026, 1, 1),
            datetime(2026, 12, 31),
        ]
    )

    # Blank row harus diabaikan.
    worksheet.append(
        [
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
        ]
    )

    file_buffer = BytesIO()

    workbook.save(file_buffer)
    workbook.close()

    result = read_portfolio_workbook(
        file_buffer.getvalue()
    )

    assert result["total_rows"] == 1
    assert result["valid_rows"] == 1
    assert result["invalid_rows"] == 0

    assert result["rows"][0]["row"] == 3
    assert result["rows"][0]["valid"] is True

    assert (
        result["rows"][0]["data"]["project_code"]
        == "PRJ-001"
    )


def test_read_portfolio_workbook_missing_sheet():
    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "WrongSheet"

    file_buffer = BytesIO()

    workbook.save(file_buffer)
    workbook.close()

    with pytest.raises(
        ValueError,
        match="Sheet 'Sheet1' tidak ditemukan",
    ):
        read_portfolio_workbook(
            file_buffer.getvalue()
        )
def test_read_portfolio_workbook_invalid_file():
    with pytest.raises(
        ValueError,
        match=(
            "File Excel tidak valid "
            "atau tidak dapat dibaca"
        ),
    ):
        read_portfolio_workbook(
            b"ini bukan file excel"
        )
def test_build_portfolio_import_template():
    file_content = build_portfolio_import_template()

    workbook = load_workbook(
        BytesIO(file_content),
        read_only=False,
    )

    worksheet = workbook["Sheet1"]

    assert (
        worksheet["A1"].value
        == "Customer Perspective"
    )
    assert (
        worksheet["I1"].value
        == "Project Perspective"
    )
    assert (
        worksheet["AC1"].value
        == "Financial Perspective"
    )
    assert (
        worksheet["AG1"].value
        == "Tech Stack"
    )
    assert (
        worksheet["AK1"].value
        == "Partner Perspective"
    )

    assert worksheet["A2"].value == "No"
    assert worksheet["B2"].value == "Customer"
    assert worksheet["O2"].value == "Project ID"
    assert (
        worksheet["P2"].value
        == "Project Value (before tax)"
    )
    assert worksheet["AK2"].value == "Partner"

    assert worksheet.max_column == 37
    assert worksheet.freeze_panes == "A3"

    workbook.close()